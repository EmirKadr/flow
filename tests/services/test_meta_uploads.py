import asyncio
import io
from pathlib import Path
import re
from types import SimpleNamespace

import pytest
from fastapi import BackgroundTasks, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from starlette.datastructures import Headers, UploadFile

from app.backend.config import Settings, settings
from app.backend.database import Base
from app.backend.deps import get_current_user, get_db
from app.backend.main import app
from app.backend.models import AuditLog, MetaMediaUpload, MetaShipmentObservation, User
from app.backend import meta_analysis_service
from app.backend.routers import meta_uploads
from app.backend.routers import meta_uploads_helpers


@pytest.fixture(autouse=True)
def disable_meta_analysis_provider(monkeypatch, tmp_path):
    from app.backend import media_store

    monkeypatch.setattr(settings, "GEMINI_API_KEY", "")
    monkeypatch.setattr(settings, "MEDIA_STORE_ROOT", str(tmp_path / "media_store"))
    monkeypatch.setattr(meta_uploads, "_probe_video_duration_from_path", lambda path: None)
    media_store.reset_media_store_cache()
    yield
    media_store.reset_media_store_cache()


def make_session():
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    session = SessionLocal()
    return engine, session


def make_upload(filename: str, content: bytes, content_type: str) -> UploadFile:
    return UploadFile(
        file=io.BytesIO(content),
        filename=filename,
        headers=Headers({"content-type": content_type}),
    )


def store_bytes(content: bytes, suffix: str = ".bin") -> str:
    """Skriv bytes till MediaStore och returnera storage_key (för testrader)."""
    from app.backend.media_store import get_media_store

    writer = get_media_store().create_writer(suffix=suffix)
    writer.write(content)
    return writer.commit().key


def test_meta_upload_default_rate_limit_allows_long_sequential_queues(monkeypatch):
    monkeypatch.delenv("META_UPLOAD_RATE_LIMIT_PER_MINUTE", raising=False)
    defaults = Settings(_env_file=None)
    assert defaults.META_UPLOAD_RATE_LIMIT_PER_MINUTE == 0
    assert defaults.MAX_META_UPLOAD_FILES >= 1


def test_meta_analysis_uses_gemini_config_and_parses_json_response(monkeypatch):
    monkeypatch.setattr(settings, "GEMINI_API_KEY", "gemini-key")
    monkeypatch.setattr(settings, "GEMINI_MODEL", "gemini-2.5-pro")

    assert meta_analysis_service.meta_analysis_configured()
    assert meta_analysis_service.gemini_model_name() == "gemini-2.5-pro"

    extracted = meta_analysis_service._extract_json_candidate(
        {
            "candidates": [
                {
                    "content": {
                        "parts": [
                            {
                                "text": (
                                    '{"order_number":"12345","username":"lots1","customer_name":"Kund AB",'
                                    '"shipment_number":"RIG-REF-260601-0000000-0",'
                                    '"pallet_id":"PALL-1","deviations":["Dåligt byggd pall"],'
                                    '"uncertainty_notes":"Kontrollera kundnamn","label_frame_time_seconds":2.5}'
                                )
                            }
                        ]
                    }
                }
            ]
        }
    )
    fields = meta_analysis_service.normalize_meta_analysis(extracted)

    assert fields["order_number"] is None
    assert fields["shipment_number"] is None
    assert fields["username"] is None
    assert fields["customer_name"] is None
    assert fields["pallet_id"] == "PALL-1"
    assert fields["deviations"] == ["Dåligt byggd pall"]
    assert fields["uncertainty_notes"] == "Kontrollera kundnamn"
    assert "label_frame_time_seconds" not in fields
    assert re.fullmatch(
        r"[0-9a-f]{64}",
        meta_analysis_service.calculate_record_hash(
            video_hash="b" * 64,
            order_number=fields["order_number"],
            shipment_number=fields["shipment_number"],
            username=fields["username"],
            customer_name=fields["customer_name"],
            pallet_id=fields["pallet_id"],
            deviations=fields["deviations"],
        ),
    )


def test_meta_analysis_normalizes_audio_only_fields():
    fields = meta_analysis_service.normalize_meta_analysis(
        {
            "order_numbers": ["100001", "100002"],
            "sändnings-id": "RIG-REF-260601-0000000-0",
            "användarnamn": "LOTS01",
            "kund": "Kund AB",
            "pallet_ids": ["BOX-001", "BOX-002"],
            "avvikelser": [{"description": "Pall behöver plastas"}],
        }
    )

    assert fields["order_number"] is None
    assert fields["shipment_number"] is None
    assert fields["username"] is None
    assert fields["customer_name"] is None
    assert fields["pallet_id"] == "BOX-001"
    assert fields["deviations"] == ["Pall behöver plastas"]


def test_meta_analysis_treats_godsmarkning_as_pallet_id():
    fields = meta_analysis_service.normalize_meta_analysis(
        {
            "godsm\u00e4rkning": "PALL-7788",
            "deviations": ["Fel m\u00e5tt"],
        }
    )

    assert fields["pallet_id"] == "PALL-7788"
    assert fields["deviations"] == ["Fel m\u00e5tt"]


def test_meta_record_hash_ignores_future_lookup_fields():
    base = {
        "video_hash": "b" * 64,
        "order_number": "100001",
        "username": "LOTS01",
        "customer_name": "Kund AB",
        "pallet_id": "BOX-001",
        "deviations": ["Pall behöver plastas"],
    }

    first = meta_analysis_service.calculate_record_hash(**base, shipment_number="RIG-REF-260601-0000000-0")
    second = meta_analysis_service.calculate_record_hash(**base, shipment_number="RIG-REF-260601-0000001-0")
    third = meta_analysis_service.calculate_record_hash(**{**base, "pallet_id": "BOX-002"})

    assert re.fullmatch(r"[0-9a-f]{64}", first)
    assert re.fullmatch(r"[0-9a-f]{64}", second)
    assert first == second
    assert first != third


def test_meta_analysis_status_requires_only_pallet_id_and_deviations():
    assert meta_analysis_service._status_for_analysis(
        {"pallet_id": "BOX-001", "deviations": ["Pall lutar"], "uncertainty_notes": None},
    ) == "analyzed"
    assert meta_analysis_service._status_for_analysis(
        {"pallet_id": None, "deviations": ["Pall lutar"], "uncertainty_notes": None},
    ) == "manual_review"
    assert meta_analysis_service._status_for_analysis(
        {"pallet_id": "BOX-001", "deviations": [], "uncertainty_notes": None},
    ) == "manual_review"
    assert meta_analysis_service._status_for_analysis(
        {"pallet_id": "BOX-001", "deviations": ["Pall lutar"], "uncertainty_notes": "Flera pall-id hors."},
    ) == "manual_review"


def test_meta_ffmpeg_resolver_uses_imageio_fallback(monkeypatch):
    monkeypatch.setattr(meta_analysis_service.shutil, "which", lambda name: None)
    monkeypatch.setattr(meta_analysis_service, "_imageio_ffmpeg_exe", lambda: "bundled-ffmpeg")

    assert meta_analysis_service._resolve_ffmpeg() == "bundled-ffmpeg"


def test_analyze_meta_upload_keeps_lookup_fields_blank_without_dispatch_api(monkeypatch):
    monkeypatch.setattr(settings, "GEMINI_API_KEY", "gemini-key")
    monkeypatch.setattr(
        meta_analysis_service,
        "_call_meta_analysis_provider",
        lambda upload: {
            "order_number": "100001",
            "shipment_number": "RIG-REF-1",
            "username": "lots1",
            "customer_name": "Kund AB",
            "pallet_id": "BOX-001",
            "deviations": ["Pall lutar"],
        },
    )
    monkeypatch.setattr(meta_analysis_service, "workflow_api_configured", lambda: False)
    engine, session = make_session()
    row = MetaMediaUpload(
        batch_id="batch",
        original_filename="voice.mov",
        stored_filename="20260531_120102_123456Z_01.mov",
        content_type="video/quicktime",
        media_type="video",
        size_bytes=11,
        content_hash="b" * 64,
        storage_backend="filesystem",
        storage_key=store_bytes(b"hello-video", ".mov"),
        status="pending_analysis",
        source="public_upload",
    )
    session.add(row)
    session.commit()
    session.refresh(row)
    try:
        result = meta_analysis_service.analyze_meta_upload(session, row.id)

        # Utan konfigurerad extern datakalla ska uppslaget bli en begriplig
        # anteckning (inte tysta tomma falt) och raden hamna i Kontrollera.
        assert result.analysis_status == "manual_review"
        assert "inte konfigurerad" in (result.uncertainty_notes or "")
        assert result.order_number is None
        assert result.shipment_number is None
        assert result.username is None
        assert result.customer_name is None
        assert result.pallet_id == "BOX-001"
        assert result.deviations == ["Pall lutar"]
        assert result.label_image_upload_id is None
        assert result.label_frame_time_seconds is None
    finally:
        session.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_analyze_meta_upload_enriches_lookup_fields_from_dispatch_api(monkeypatch):
    monkeypatch.setattr(settings, "GEMINI_API_KEY", "gemini-key")
    monkeypatch.setattr(
        meta_analysis_service,
        "_call_meta_analysis_provider",
        lambda upload: {
            "pallet_id": "BOX-001",
            "deviations": ["Pall lutar"],
        },
    )
    monkeypatch.setattr(meta_analysis_service, "workflow_api_configured", lambda: True)

    class FakeDispatchClient:
        def __init__(self):
            self.calls = []

        def fetch_data(self, view, filters=None, identifiers=None):
            self.calls.append((view, filters, identifiers))
            return [
                {
                    "pick_pall_num": "BOX-001",
                    "order_num": "100001",
                    "shipment_id": "RIG-REF-1",
                    "user_id": "LOTS01",
                    "custom_desc": "Kund AB",
                }
            ]

    fake_client = FakeDispatchClient()
    monkeypatch.setattr(meta_analysis_service, "_api_client", lambda: fake_client)
    engine, session = make_session()
    row = MetaMediaUpload(
        batch_id="batch",
        original_filename="voice.mov",
        stored_filename="20260531_120102_123456Z_01.mov",
        content_type="video/quicktime",
        media_type="video",
        size_bytes=11,
        content_hash="b" * 64,
        storage_backend="filesystem",
        storage_key=store_bytes(b"hello-video", ".mov"),
        status="pending_analysis",
        source="public_upload",
    )
    session.add(row)
    session.commit()
    session.refresh(row)
    try:
        result = meta_analysis_service.analyze_meta_upload(session, row.id)

        assert fake_client.calls == [
            (
                "v_ask_dispatch_pallet",
                [{"id": "pick_pall_num", "value": "BOX-001", "operator": "EQ"}],
                None,
            )
        ]
        assert result.analysis_status == "analyzed"
        assert result.order_number == "100001"
        assert result.shipment_number == "RIG-REF-1"
        assert result.username == "LOTS01"
        assert result.customer_name == "Kund AB"
        assert result.pallet_id == "BOX-001"
        assert result.deviations == ["Pall lutar"]
    finally:
        session.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_analyze_meta_upload_marks_audio_extraction_failure(monkeypatch):
    monkeypatch.setattr(settings, "GEMINI_API_KEY", "gemini-key")
    monkeypatch.setattr(meta_analysis_service, "_resolve_ffmpeg", lambda: None)
    engine, session = make_session()
    row = MetaMediaUpload(
        batch_id="batch",
        original_filename="silent.mov",
        stored_filename="20260531_120102_123456Z_01.mov",
        content_type="video/quicktime",
        media_type="video",
        size_bytes=11,
        content_hash="b" * 64,
        storage_backend="filesystem",
        storage_key=store_bytes(b"hello-video", ".mov"),
        status="pending_analysis",
        source="public_upload",
    )
    session.add(row)
    session.commit()
    session.refresh(row)
    try:
        result = meta_analysis_service.analyze_meta_upload(session, row.id)

        assert result.analysis_status == "analysis_failed"
        assert "ffmpeg saknas" in (result.analysis_error or "")
    finally:
        session.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_analyze_meta_upload_marks_unexpected_error_instead_of_leaking_500(monkeypatch):
    """Ett oväntat fel (inte MetaAnalysisFailed) får inte lämna raden låst i
    'analyzing' med en naken 500 — den ska bli 'analysis_failed' med felet synligt
    så Analysera-knappen inte disablas för alltid."""
    monkeypatch.setattr(settings, "GEMINI_API_KEY", "gemini-key")

    def _boom(upload):
        raise RuntimeError("oväntad krasch")

    monkeypatch.setattr(meta_analysis_service, "_call_meta_analysis_provider", _boom)
    engine, session = make_session()
    row = MetaMediaUpload(
        batch_id="batch",
        original_filename="voice.mov",
        stored_filename="20260531_120102_123456Z_01.mov",
        content_type="video/quicktime",
        media_type="video",
        size_bytes=11,
        content_hash="c" * 64,
        storage_backend="filesystem",
        storage_key=store_bytes(b"hello-video", ".mov"),
        status="pending_analysis",
        source="public_upload",
    )
    session.add(row)
    session.commit()
    session.refresh(row)
    try:
        result = meta_analysis_service.analyze_meta_upload(session, row.id)

        assert result.analysis_status == "analysis_failed"
        assert "RuntimeError" in (result.analysis_error or "")
        assert "oväntad krasch" in (result.analysis_error or "")
    finally:
        session.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_analyze_endpoint_returns_200_not_500_on_unexpected_error(monkeypatch):
    """Endpoint-nivå: ett oväntat fel i analysen får inte läcka som HTTP 500. Med
    buggen propagerade undantaget (TestClient reser det); med fixen svarar
    endpointen 200 med analysis_failed och felet i message, och raden är inte
    kvar i 'analyzing' (Analysera-knappen disablas annars för alltid)."""
    monkeypatch.setattr(settings, "GEMINI_API_KEY", "gemini-key")

    def _boom(upload):
        raise RuntimeError("oväntad krasch")

    monkeypatch.setattr(meta_analysis_service, "_call_meta_analysis_provider", _boom)
    engine, session = make_session()
    row = MetaMediaUpload(
        batch_id="batch-500",
        original_filename="voice.mov",
        stored_filename="20260531_120102_123456Z_01.mov",
        content_type="video/quicktime",
        media_type="video",
        size_bytes=11,
        content_hash="e" * 64,
        storage_backend="filesystem",
        storage_key=store_bytes(b"boom-video", ".mov"),
        status="pending_analysis",
        source="public_upload",
    )
    session.add(row)
    session.commit()
    session.refresh(row)
    upload_id = row.id

    def override_get_db():
        yield session

    def super_user():
        return User(id=99, username="root", role="super_user", roles=["super_user"], is_active=True)

    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[get_current_user] = super_user
    try:
        client = TestClient(app)
        response = client.post(f"/api/meta/uploads/{upload_id}/analyze")

        assert response.status_code == 200
        body = response.json()
        assert body["status"] == "analysis_failed"
        assert "RuntimeError" in (body["message"] or "")
        observation = (
            session.query(MetaShipmentObservation).filter_by(media_upload_id=upload_id).first()
        )
        assert observation is not None
        assert observation.analysis_status == "analysis_failed"
    finally:
        app.dependency_overrides.pop(get_db, None)
        app.dependency_overrides.pop(get_current_user, None)
        session.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_queued_meta_analysis_upload_ids_returns_oldest_queued_video():
    engine, session = make_session()
    first = MetaMediaUpload(
        batch_id="batch",
        original_filename="first.mov",
        stored_filename="first.mov",
        content_type="video/quicktime",
        media_type="video",
        size_bytes=11,
        content_hash="1" * 64,
        storage_backend="filesystem",
        storage_key=store_bytes(b"first-video", ".mov"),
        status="pending_analysis",
        source="public_upload",
    )
    second = MetaMediaUpload(
        batch_id="batch",
        original_filename="second.mov",
        stored_filename="second.mov",
        content_type="video/quicktime",
        media_type="video",
        size_bytes=11,
        content_hash="2" * 64,
        storage_backend="filesystem",
        storage_key=store_bytes(b"second-video", ".mov"),
        status="pending_analysis",
        source="public_upload",
    )
    session.add_all([first, second])
    session.commit()
    session.refresh(first)
    session.refresh(second)
    session.add_all(
        [
            MetaShipmentObservation(
                media_upload_id=first.id,
                video_hash=first.content_hash,
                record_hash="a" * 64,
                analysis_status="queued",
            ),
            MetaShipmentObservation(
                media_upload_id=second.id,
                video_hash=second.content_hash,
                record_hash="b" * 64,
                analysis_status="manual_review",
            ),
        ]
    )
    session.commit()
    try:
        assert meta_analysis_service.queued_meta_analysis_upload_ids(session, limit=5) == [first.id]
    finally:
        session.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_public_meta_upload_route_accepts_multiple_media_without_login(monkeypatch):
    monkeypatch.setattr(meta_uploads, "_probe_video_duration_from_path", lambda path: 42.4)
    engine, session = make_session()

    def override_get_db():
        yield session

    app.dependency_overrides[get_db] = override_get_db
    try:
        client = TestClient(app)
        response = client.post(
            "/api/meta/uploads",
            files=[
                ("files", ("bild.jpg", b"image-bytes", "image/jpeg")),
                ("files", ("film.mov", b"video-bytes", "video/quicktime")),
            ],
        )

        assert response.status_code == 201
        payload = response.json()
        assert payload["saved_count"] == 2
        assert payload["status"] == "pending_analysis"

        rows = session.query(MetaMediaUpload).order_by(MetaMediaUpload.id).all()
        assert [row.original_filename for row in rows] == ["bild.jpg", "film.mov"]
        assert re.fullmatch(r"\d{8}_\d{6}_\d{6}Z_01\.jpg", rows[0].stored_filename)
        assert re.fullmatch(r"\d{8}_\d{6}_\d{6}Z_02\.mov", rows[1].stored_filename)
        assert [item["filename"] for item in payload["saved"]] == [row.stored_filename for row in rows]
        assert all(re.fullmatch(r"[0-9a-f]{64}", row.content_hash or "") for row in rows)
        assert len({row.content_hash for row in rows}) == 2
        assert [row.media_type for row in rows] == ["image", "video"]
        # Bytena ligger i MediaStore, inte i DB-kolumnen.
        from app.backend.media_store import get_media_store

        store = get_media_store()
        assert rows[0].storage_key and rows[1].storage_key
        assert rows[0].storage_backend == "filesystem"
        assert b"".join(store.open_all(rows[0].storage_key)) == b"image-bytes"
        assert b"".join(store.open_all(rows[1].storage_key)) == b"video-bytes"
        assert rows[0].storage_key and rows[1].storage_key
        assert rows[0].duration_seconds is None
        assert rows[1].duration_seconds == 42.4
        assert len({row.batch_id for row in rows}) == 1
        assert payload["saved"][1]["duration_seconds"] == 42.4
        assert payload["saved"][1]["duration_label"] == "0:42"
        shipment = session.query(MetaShipmentObservation).one()
        assert shipment.media_upload_id == rows[1].id
        assert shipment.video_hash == rows[1].content_hash
        assert re.fullmatch(r"[0-9a-f]{64}", shipment.record_hash)
        assert shipment.analysis_status == "needs_configuration"
        audit = session.query(AuditLog).filter_by(entity_type="meta_media_upload", action="upload_success").one()
        assert audit.user_id is None
        assert audit.new_value["status_code"] == 201
        assert audit.new_value["attempted_count"] == 2
        assert audit.new_value["saved_count"] == 2
        assert audit.new_value["skipped_count"] == 0
        assert audit.new_value["shipment_count"] == 1
        assert audit.new_value["uploaded_bytes"] == len(b"image-bytes") + len(b"video-bytes")
        assert audit.new_value["media_types"] == ["image", "video"]
        assert audit.new_value["analysis_status"] == "needs_configuration"
        assert "bild.jpg" not in str(audit.new_value)
        assert "film.mov" not in str(audit.new_value)
    finally:
        app.dependency_overrides.pop(get_db, None)
        session.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_public_meta_upload_server_failure_is_visible_in_error_audit(monkeypatch):
    engine, session = make_session()
    original_commit = session.commit
    commit_calls = 0

    def flaky_commit():
        nonlocal commit_calls
        commit_calls += 1
        if commit_calls == 1:
            raise RuntimeError("database write failed")
        return original_commit()

    monkeypatch.setattr(session, "commit", flaky_commit)

    def override_get_db():
        yield session

    def super_user():
        return User(id=99, username="root", role="super_user", roles=["super_user"], is_active=True)

    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[get_current_user] = super_user
    try:
        client = TestClient(app)
        response = client.post(
            "/api/meta/uploads",
            files=[("files", ("film.mp4", b"video-bytes", "video/mp4"))],
        )

        assert response.status_code == 500
        assert "Uppladdningen misslyckades" in response.json()["detail"]
        assert session.query(MetaMediaUpload).count() == 0

        audit = session.query(AuditLog).filter_by(entity_type="meta_media_upload", action="upload_failed").one()
        assert audit.user_id is None
        assert audit.new_value["status_code"] == 500
        assert audit.new_value["error_code"] == "HTTP 500"
        assert audit.new_value["error_type"] == "RuntimeError"
        assert audit.new_value["path"] == "/api/meta/uploads"
        assert audit.new_value["attempted_count"] == 1
        assert audit.new_value["accepted_count"] == 1
        assert audit.new_value["uploaded_bytes"] == len(b"video-bytes")
        assert "film.mp4" not in str(audit.new_value)
        assert "database write failed" not in str(audit.new_value)

        errors = client.get("/api/audit/errors")
        assert errors.status_code == 200
        event = errors.json()["recent"][0]
        assert event["entity_type"] == "meta_media_upload"
        assert event["action"] == "upload_failed"
        assert event["status_code"] == 500
        assert event["path"] == "/api/meta/uploads"
        assert event["username"] is None
    finally:
        app.dependency_overrides.pop(get_db, None)
        app.dependency_overrides.pop(get_current_user, None)
        session.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_meta_upload_skips_duplicate_media_bytes():
    engine, session = make_session()

    def override_get_db():
        yield session

    app.dependency_overrides[get_db] = override_get_db
    try:
        client = TestClient(app)
        first = client.post(
            "/api/meta/uploads",
            files=[
                ("files", ("film-a.mov", b"same-video", "video/quicktime")),
                ("files", ("film-b.mov", b"same-video", "video/quicktime")),
            ],
        )

        assert first.status_code == 201
        first_payload = first.json()
        assert first_payload["saved_count"] == 1
        assert first_payload["skipped_count"] == 1
        assert first_payload["skipped"][0]["reason"] == "duplicate"
        assert first_payload["skipped"][0]["duplicate_of_filename"] == first_payload["saved"][0]["filename"]
        assert session.query(MetaMediaUpload).count() == 1

        second = client.post(
            "/api/meta/uploads",
            files=[("files", ("film-c.mov", b"same-video", "video/quicktime"))],
        )
        assert second.status_code == 201
        second_payload = second.json()
        assert second_payload["saved_count"] == 0
        assert second_payload["skipped_count"] == 1
        assert second_payload["skipped"][0]["duplicate_of_id"] is not None
        assert session.query(MetaMediaUpload).count() == 1
    finally:
        app.dependency_overrides.pop(get_db, None)
        session.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_meta_upload_rejects_non_media_files():
    engine, session = make_session()
    fake_request = SimpleNamespace(client=SimpleNamespace(host="test-client"))
    try:
        with pytest.raises(HTTPException) as exc_info:
            asyncio.run(
                meta_uploads.upload_meta_media(
                    request=fake_request,
                    background_tasks=BackgroundTasks(),
                    files=[make_upload("anteckning.txt", b"text", "text/plain")],
                    db=session,
                )
            )

        assert exc_info.value.status_code == 400
        assert session.query(MetaMediaUpload).count() == 0
    finally:
        session.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_super_user_can_list_meta_uploads_and_stream_content():
    engine, session = make_session()
    row = MetaMediaUpload(
        batch_id="batch",
        original_filename="semesterfilm.mov",
        stored_filename="20260531_120102_123456Z_01.mov",
        content_type="video/quicktime",
        media_type="video",
        size_bytes=11,
        duration_seconds=65.1,
        storage_backend="filesystem",
        storage_key=store_bytes(b"hello-video", ".mov"),
        status="pending_analysis",
        source="public_upload",
    )
    session.add(row)
    session.commit()
    session.refresh(row)

    def override_get_db():
        yield session

    def super_user():
        return User(id=99, username="root", role="super_user", roles=["super_user"], is_active=True)

    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[get_current_user] = super_user
    try:
        client = TestClient(app)
        response = client.get("/api/meta/uploads")
        assert response.status_code == 200
        item = response.json()["items"][0]
        assert item["filename"] == "20260531_120102_123456Z_01.mov"
        assert item["original_filename"] == "semesterfilm.mov"
        assert item["media_type"] == "video"
        assert item["duration_seconds"] == 65.1
        assert item["duration_label"] == "1:05"

        content = client.get(f"/api/meta/uploads/{row.id}/content", headers={"Range": "bytes=0-4"})
        assert content.status_code == 206
        assert content.content == b"hello"
        assert content.headers["content-range"] == "bytes 0-4/11"
        assert content.headers["content-disposition"].startswith("inline;")
        assert "20260531_120102_123456Z_01.mov" in content.headers["content-disposition"]

        head = client.head(f"/api/meta/uploads/{row.id}/content?download=1")
        assert head.status_code == 200
        assert head.content == b""
        assert head.headers["content-type"].startswith("video/quicktime")
        assert head.headers["content-length"] == "11"
        assert head.headers["content-disposition"].startswith("attachment;")

        download = client.get(f"/api/meta/uploads/{row.id}/content?download=1", headers={"Range": "bytes=0-4"})
        assert download.status_code == 206
        assert download.content == b"hello"
        assert download.headers["content-disposition"].startswith("attachment;")
    finally:
        app.dependency_overrides.pop(get_db, None)
        app.dependency_overrides.pop(get_current_user, None)
        session.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_meta_video_content_response_normalizes_audio_mime_from_upload_client():
    engine, session = make_session()
    row = MetaMediaUpload(
        batch_id="batch-audio-mime",
        original_filename="metavideo.mp4",
        stored_filename="20260531_120102_123456Z_01.mp4",
        content_type="audio/mp4",
        media_type="video",
        size_bytes=11,
        duration_seconds=65.1,
        storage_backend="filesystem",
        storage_key=store_bytes(b"hello-video", ".mp4"),
        status="pending_analysis",
        source="public_upload",
    )
    session.add(row)
    session.commit()
    session.refresh(row)

    def override_get_db():
        yield session

    def super_user():
        return User(id=99, username="root", role="super_user", roles=["super_user"], is_active=True)

    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[get_current_user] = super_user
    try:
        client = TestClient(app)
        content = client.get(f"/api/meta/uploads/{row.id}/content")

        assert content.status_code == 200
        assert content.content == b"hello-video"
        assert content.headers["content-type"].startswith("video/mp4")
        assert "20260531_120102_123456Z_01.mp4" in content.headers["content-disposition"]

        ranged = client.get(f"/api/meta/uploads/{row.id}/content", headers={"Range": "bytes=0-4"})
        assert ranged.status_code == 206
        assert ranged.headers["content-type"].startswith("video/mp4")
    finally:
        app.dependency_overrides.pop(get_db, None)
        app.dependency_overrides.pop(get_current_user, None)
        session.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_public_meta_upload_normalizes_audio_mp4_content_type_for_video_filename():
    engine, session = make_session()

    def override_get_db():
        yield session

    app.dependency_overrides[get_db] = override_get_db
    try:
        client = TestClient(app)
        response = client.post(
            "/api/meta/uploads",
            files=[("files", ("film.mp4", b"video-bytes", "audio/mp4"))],
        )

        assert response.status_code == 201
        row = session.query(MetaMediaUpload).one()
        assert row.media_type == "video"
        assert row.content_type == "video/mp4"
        assert response.json()["saved"][0]["content_type"] == "video/mp4"
    finally:
        app.dependency_overrides.pop(get_db, None)
        session.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_meta_video_content_playable_variant_transcodes_to_h264_download(monkeypatch):
    engine, session = make_session()
    row = MetaMediaUpload(
        batch_id="batch-playable",
        original_filename="metavideo.mp4",
        stored_filename="20260531_120102_123456Z_01.mp4",
        content_type="video/mp4",
        media_type="video",
        size_bytes=11,
        duration_seconds=65.1,
        storage_backend="filesystem",
        storage_key=store_bytes(b"source-video", ".mp4"),
        status="pending_analysis",
        source="public_upload",
    )
    session.add(row)
    session.commit()
    session.refresh(row)

    def fake_run(command, capture_output, timeout, check):
        Path(command[-1]).write_bytes(b"playable-video")
        return SimpleNamespace(returncode=0, stderr=b"")

    monkeypatch.setattr(meta_uploads_helpers, "_resolve_ffmpeg", lambda: "ffmpeg")
    monkeypatch.setattr(meta_uploads_helpers.subprocess, "run", fake_run)
    queue_events = []

    class FakeSemaphore:
        def __enter__(self):
            queue_events.append("enter")

        def __exit__(self, exc_type, exc, tb):
            queue_events.append("exit")

    monkeypatch.setattr(meta_uploads_helpers, "_PLAYABLE_TRANSCODE_SEMAPHORE", FakeSemaphore())

    def override_get_db():
        yield session

    def super_user():
        return User(id=99, username="root", role="super_user", roles=["super_user"], is_active=True)

    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[get_current_user] = super_user
    try:
        client = TestClient(app)
        content = client.get(f"/api/meta/uploads/{row.id}/content?variant=playable")

        assert content.status_code == 200
        assert content.content == b"playable-video"
        assert content.headers["content-type"].startswith("video/mp4")
        assert "20260531_120102_123456Z_01.mp4" in content.headers["content-disposition"]
        assert queue_events == ["enter", "exit"]
    finally:
        app.dependency_overrides.pop(get_db, None)
        app.dependency_overrides.pop(get_current_user, None)
        session.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_super_user_can_list_and_request_meta_shipment_analysis_without_configuration(monkeypatch):
    monkeypatch.setattr(meta_uploads, "meta_analysis_configured", lambda: False)
    engine, session = make_session()
    row = MetaMediaUpload(
        batch_id="batch",
        original_filename="etikettfilm.mov",
        stored_filename="20260531_120102_123456Z_01.mov",
        content_type="video/quicktime",
        media_type="video",
        size_bytes=11,
        duration_seconds=75.4,
        content_hash="b" * 64,
        storage_backend="filesystem",
        storage_key=store_bytes(b"hello-video", ".mov"),
        status="pending_analysis",
        source="public_upload",
    )
    session.add(row)
    session.commit()
    session.refresh(row)
    shipment = MetaShipmentObservation(
        media_upload_id=row.id,
        video_hash=row.content_hash,
        record_hash="c" * 64,
        order_number=None,
        shipment_number=None,
        username=None,
        customer_name=None,
        pallet_id="PALL-1",
        deviations=["Dåligt byggd pall"],
        analysis_status="manual_review",
    )
    session.add(shipment)
    session.commit()

    def override_get_db():
        yield session

    def super_user():
        return User(id=99, username="root", role="super_user", roles=["super_user"], is_active=True)

    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[get_current_user] = super_user
    try:
        client = TestClient(app)
        response = client.get("/api/meta/shipment-observations")
        assert response.status_code == 200
        item = response.json()["items"][0]
        assert item["order_number"] is None
        assert item["shipment_number"] is None
        assert item["username"] is None
        assert item["customer_name"] is None
        assert item["pallet_id"] == "PALL-1"
        assert item["deviations"] == ["Dåligt byggd pall"]
        assert item["video_url"] == f"/api/meta/uploads/{row.id}/content"
        assert item["video_filename"] == "20260531_120102_123456Z_01.mov"
        assert item["video_duration_seconds"] == 75.4
        assert item["video_duration_label"] == "1:15"
        assert item["video_size_bytes"] == 11
        assert item["video_size_label"] == "11 B"
        assert item["created_at"]
        assert item["updated_at"]

        analysis = client.post(f"/api/meta/uploads/{row.id}/analyze")
        assert analysis.status_code == 200
        assert analysis.json()["status"] == "needs_configuration"
    finally:
        app.dependency_overrides.pop(get_db, None)
        app.dependency_overrides.pop(get_current_user, None)
        session.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_super_user_can_export_meta_shipment_observations_excel():
    engine, session = make_session()
    row = MetaMediaUpload(
        batch_id="batch-export",
        original_filename="etikettfilm.mov",
        stored_filename="20260531_120102_123456Z_01.mov",
        content_type="video/quicktime",
        media_type="video",
        size_bytes=11,
        duration_seconds=75.4,
        content_hash="b" * 64,
        storage_backend="filesystem",
        storage_key=store_bytes(b"hello-video", ".mov"),
        status="pending_analysis",
        source="public_upload",
    )
    session.add(row)
    session.commit()
    session.refresh(row)
    shipment = MetaShipmentObservation(
        media_upload_id=row.id,
        video_hash=row.content_hash,
        record_hash="c" * 64,
        order_number="100001",
        shipment_number="RIG-REF-1",
        username="LOTS01",
        customer_name="Kund AB",
        pallet_id="PALL-1",
        deviations=["D\u00e5ligt byggd pall"],
        analysis_status="manual_review",
        uncertainty_notes="Kontrollera",
    )
    session.add(shipment)
    session.commit()
    session.refresh(shipment)

    def override_get_db():
        yield session

    def super_user():
        return User(id=99, username="root", role="super_user", roles=["super_user"], is_active=True)

    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[get_current_user] = super_user
    try:
        client = TestClient(app)
        response = client.get(f"/api/meta/shipment-observations/export?ids={shipment.id}")

        assert response.status_code == 200
        assert response.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        workbook = load_workbook(io.BytesIO(response.content), read_only=True, data_only=True)
        sheet = workbook.active
        assert [cell.value for cell in sheet[1]][:7] == [
            "Ordernummer",
            "S\u00e4ndningsnummer",
            "Anv\u00e4ndarnamn",
            "Kund",
            "Pall-ID",
            "Avvikelser",
            "Status",
        ]
        assert [cell.value for cell in sheet[1]][9:12] == ["Video", "L\u00e4ngd", "Storlek"]
        values = [cell.value for cell in sheet[2]]
        assert values[:8] == [
            "100001",
            "RIG-REF-1",
            "LOTS01",
            "Kund AB",
            "PALL-1",
            "D\u00e5ligt byggd pall",
            "manual_review",
            "Kontrollera",
        ]
        assert values[9] == "20260531_120102_123456Z_01.mov"
        assert values[10] == "1:15"
        assert values[11] == "11 B"
        assert values[12] == "c" * 64
        workbook.close()
    finally:
        app.dependency_overrides.pop(get_db, None)
        app.dependency_overrides.pop(get_current_user, None)
        session.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_super_user_can_delete_meta_uploads_and_audit_without_blob():
    engine, session = make_session()
    row = MetaMediaUpload(
        batch_id="batch-delete",
        original_filename="lagerbild.jpg",
        stored_filename="20260531_120102_123456Z_01.jpg",
        content_type="image/jpeg",
        media_type="image",
        size_bytes=10,
        content_hash="a" * 64,
        storage_backend="filesystem",
        storage_key=store_bytes(b"image-data", ".jpg"),
        status="pending_analysis",
        source="public_upload",
    )
    session.add(row)
    session.commit()
    session.refresh(row)
    upload_id = row.id

    def override_get_db():
        yield session

    def super_user():
        return User(id=99, username="root", role="super_user", roles=["super_user"], is_active=True)

    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[get_current_user] = super_user
    try:
        client = TestClient(app)
        response = client.delete(f"/api/meta/uploads/{upload_id}")

        assert response.status_code == 204
        assert session.get(MetaMediaUpload, upload_id) is None
        audit = session.query(AuditLog).filter_by(entity_type="meta_media_upload", action="delete").one()
        assert audit.entity_id == upload_id
        assert audit.old_value["filename"] == "20260531_120102_123456Z_01.jpg"
        assert audit.old_value["content_hash"] == "a" * 64
        assert "data" not in audit.old_value
    finally:
        app.dependency_overrides.pop(get_db, None)
        app.dependency_overrides.pop(get_current_user, None)
        session.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_super_user_can_change_shipment_status_and_audit():
    engine, session = make_session()
    row = MetaMediaUpload(
        batch_id="batch-status",
        original_filename="etikettfilm.mov",
        stored_filename="20260531_120102_123456Z_01.mov",
        content_type="video/quicktime",
        media_type="video",
        size_bytes=11,
        content_hash="d" * 64,
        storage_backend="filesystem",
        storage_key=store_bytes(b"status-video", ".mov"),
        status="pending_analysis",
        source="public_upload",
    )
    session.add(row)
    session.commit()
    session.refresh(row)
    shipment = MetaShipmentObservation(
        media_upload_id=row.id,
        video_hash=row.content_hash,
        record_hash="e" * 64,
        pallet_id="PALL-9",
        deviations=[],
        analysis_status="manual_review",
    )
    session.add(shipment)
    session.commit()
    session.refresh(shipment)
    observation_id = shipment.id

    def override_get_db():
        yield session

    def super_user():
        return User(id=99, username="root", role="super_user", roles=["super_user"], is_active=True)

    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[get_current_user] = super_user
    try:
        client = TestClient(app)
        response = client.patch(
            f"/api/meta/shipment-observations/{observation_id}/status",
            json={"status": "analyzed"},
        )
        assert response.status_code == 200
        assert response.json()["status"] == "analyzed"
        assert response.json()["item"]["analysis_status"] == "analyzed"
        session.expire_all()
        assert session.get(MetaShipmentObservation, observation_id).analysis_status == "analyzed"
        audit = session.query(AuditLog).filter_by(
            entity_type="meta_shipment_observation", action="update_status"
        ).one()
        assert audit.entity_id == observation_id
        assert audit.old_value["analysis_status"] == "manual_review"
        assert audit.new_value["analysis_status"] == "analyzed"

        rejected = client.patch(
            f"/api/meta/shipment-observations/{observation_id}/status",
            json={"status": "not-a-real-status"},
        )
        assert rejected.status_code == 400
    finally:
        app.dependency_overrides.pop(get_db, None)
        app.dependency_overrides.pop(get_current_user, None)
        session.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_non_super_user_cannot_list_meta_uploads(monkeypatch):
    monkeypatch.setattr(settings, "SUPER_USER_USERNAMES", "admin,emikad")
    engine, session = make_session()

    def override_get_db():
        yield session

    def admin_user():
        return User(id=100, username="regular-admin", role="admin", roles=["admin"], is_active=True)

    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[get_current_user] = admin_user
    try:
        client = TestClient(app)
        response = client.get("/api/meta/uploads")
        assert response.status_code == 403
        delete_response = client.delete("/api/meta/uploads/1")
        assert delete_response.status_code == 403
    finally:
        app.dependency_overrides.pop(get_db, None)
        app.dependency_overrides.pop(get_current_user, None)
        session.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_settings_blank_unsubstituted_octopus_placeholders():
    """Octopus lämnar okända #{VAR}-platshållare ordagrant kvar i manifestet.
    Ett sådant värde ska behandlas som osatt (hände 2026-07-09: varje
    meta-videoanalys kraschade med ValueError på GEMINI_API_BASE_URL)."""
    cfg = Settings(
        _env_file=None,
        GEMINI_API_BASE_URL="#{GEMINI_API_BASE_URL}",
        GEMINI_API_KEY="#{GEMINI_API_KEY}",
        MINIMAX_API_KEY=" #{MINIMAX_API_KEY} ",
    )
    assert cfg.GEMINI_API_BASE_URL == ""
    assert cfg.GEMINI_API_KEY == ""
    assert cfg.MINIMAX_API_KEY == ""
    # Riktiga värden passerar orörda.
    kept = Settings(_env_file=None, GEMINI_API_BASE_URL="https://proxy.example.com", GEMINI_API_KEY="riktig-nyckel")
    assert kept.GEMINI_API_BASE_URL == "https://proxy.example.com"
    assert kept.GEMINI_API_KEY == "riktig-nyckel"


def test_gemini_base_url_falls_back_on_non_http_values(monkeypatch):
    default = "https://generativelanguage.googleapis.com"
    monkeypatch.setattr(settings, "GEMINI_API_BASE_URL", "#{GEMINI_API_BASE_URL}")
    assert meta_analysis_service._gemini_base_url() == default
    monkeypatch.setattr(settings, "GEMINI_API_BASE_URL", "")
    assert meta_analysis_service._gemini_base_url() == default
    monkeypatch.setattr(settings, "GEMINI_API_BASE_URL", "https://proxy.example.com/")
    assert meta_analysis_service._gemini_base_url() == "https://proxy.example.com"


def test_gemini_requests_put_api_key_in_header_not_url(monkeypatch):
    """Nyckeln får aldrig ligga i URL:en — URL:er hamnar i undantagstexter,
    spans och loggar (2026-07-09 läckte ?key=... via ValueError i analysis_error)."""
    monkeypatch.setattr(settings, "GEMINI_API_KEY", "super-hemlig-nyckel")
    url = meta_analysis_service._gemini_url("/upload/v1beta/files")
    assert "super-hemlig-nyckel" not in url
    assert "key=" not in url
    headers = meta_analysis_service._gemini_headers({"Content-Type": "application/json"})
    assert headers["x-goog-api-key"] == "super-hemlig-nyckel"
    assert headers["Content-Type"] == "application/json"


def test_ffmpeg_commands_cap_threads_to_protect_pod_memory(monkeypatch):
    """ffmpeg default ar tradar for alla karnor; multitradad avkodning av
    hogupplost h264 (Meta-glasogon: 1488x1984@120fps) tog ~254 MB och
    grupp-OOM-dodade podden 2026-07-09. Ljudextraktionen ska kora
    -threads 1 och playable-transkodningen -threads 2 (avkodare + x264)."""
    captured: list[list[str]] = []

    def fake_run(command, capture_output, timeout, check):
        captured.append([str(part) for part in command])
        Path(command[-1]).write_bytes(b"fake-output")
        return SimpleNamespace(returncode=0, stderr=b"")

    upload = MetaMediaUpload(
        batch_id="batch-threads",
        original_filename="meta.mp4",
        stored_filename="20260709_010203_000000Z_01.mp4",
        content_type="video/mp4",
        media_type="video",
        size_bytes=11,
        content_hash="a1" * 32,
        storage_backend="filesystem",
        storage_key=store_bytes(b"video-bytes", ".mp4"),
        status="pending_analysis",
        source="public_upload",
    )

    monkeypatch.setattr(meta_analysis_service, "_resolve_ffmpeg", lambda: "ffmpeg")
    monkeypatch.setattr(meta_analysis_service.subprocess, "run", fake_run)
    with meta_analysis_service.extract_audio_file(upload) as audio:
        assert audio.size_bytes > 0

    monkeypatch.setattr(meta_uploads_helpers, "_resolve_ffmpeg", lambda: "ffmpeg")
    monkeypatch.setattr(meta_uploads_helpers.subprocess, "run", fake_run)
    playable = meta_uploads_helpers._transcode_video_to_playable(upload)
    playable.unlink(missing_ok=True)

    audio_cmd, transcode_cmd = captured
    index = audio_cmd.index("-threads")
    assert audio_cmd[index + 1] == "1"
    assert index < audio_cmd.index("-i"), "trådtaket ska sättas före -i (avkodaren)"
    assert transcode_cmd.count("-threads") == 2
    first = transcode_cmd.index("-threads")
    second = transcode_cmd.index("-threads", first + 1)
    assert transcode_cmd[first + 1] == "2" and transcode_cmd[second + 1] == "2"
    assert first < transcode_cmd.index("-i") < transcode_cmd.index("libx264") < second


def test_analysis_no_longer_creates_label_stills():
    """Stillbildsfunktionen ar borttagen (beslut 2026-07-09): analysen far
    aldrig ateruppliva videoavkodning i podden. Skyddar mot regression via
    ateranvanda hjalpfunktioner."""
    assert not hasattr(meta_analysis_service, "extract_label_still_bytes")
    assert not hasattr(meta_analysis_service, "create_label_still_upload")
    assert not hasattr(Settings(_env_file=None), "META_LABEL_STILL_TIME_SECONDS")


def _make_analyzed_upload(session, content_hash: str) -> MetaMediaUpload:
    row = MetaMediaUpload(
        batch_id="batch-lookup",
        original_filename="voice.mov",
        stored_filename="20260531_120102_123456Z_01.mov",
        content_type="video/quicktime",
        media_type="video",
        size_bytes=11,
        content_hash=content_hash,
        storage_backend="filesystem",
        storage_key=store_bytes(b"lookup-video" + content_hash.encode(), ".mov"),
        status="pending_analysis",
        source="public_upload",
    )
    session.add(row)
    session.commit()
    session.refresh(row)
    return row


def test_dispatch_lookup_miss_becomes_note_not_failure(monkeypatch):
    """Ingen traff i Dispatchpallar ska ge en tydlig anteckning med pall-id
    och status Kontrollera — inte ett fel."""
    monkeypatch.setattr(settings, "GEMINI_API_KEY", "gemini-key")
    monkeypatch.setattr(
        meta_analysis_service,
        "_call_meta_analysis_provider",
        lambda upload: {"pallet_id": "BOX-SAKNAS", "deviations": ["Pall lutar"]},
    )
    monkeypatch.setattr(meta_analysis_service, "workflow_api_configured", lambda: True)
    monkeypatch.setattr(
        meta_analysis_service,
        "_api_client",
        lambda: SimpleNamespace(fetch_data=lambda view, filters=None, identifiers=None: []),
    )
    engine, session = make_session()
    row = _make_analyzed_upload(session, "d1" * 32)
    try:
        result = meta_analysis_service.analyze_meta_upload(session, row.id)

        assert result.analysis_status == "manual_review"
        assert "ingen traff for pall-id BOX-SAKNAS" in (result.uncertainty_notes or "")
        assert result.analysis_error is None
    finally:
        session.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_dispatch_lookup_falls_back_to_archive_view(monkeypatch):
    """Live-vyn haller bara ~14 dagar — aldre pallar ska hittas i arkivet
    (dblog_dispatch_pallet_log) och fylla lookup-falten utan anteckning.
    Numeriska pall-id skickas som tal sa EQ matchar den numeriska kolumnen."""
    monkeypatch.setattr(settings, "GEMINI_API_KEY", "gemini-key")
    monkeypatch.setattr(
        meta_analysis_service,
        "_call_meta_analysis_provider",
        lambda upload: {"pallet_id": "8473877", "deviations": ["Fel pa kartongerna"]},
    )
    monkeypatch.setattr(meta_analysis_service, "workflow_api_configured", lambda: True)

    calls = []

    def fetch_data(view, filters=None, identifiers=None):
        calls.append((view, filters))
        if view == "dblog_dispatch_pallet_log":
            return [
                {
                    "pick_pall_num": 8473877,
                    "order_num": "PR100485115",
                    "shipment_id": "GG-404-260610-PR100485115-0",
                    "user_id": "LOTS01",
                    "custom_num": "10231",
                }
            ]
        return []

    monkeypatch.setattr(meta_analysis_service, "_api_client", lambda: SimpleNamespace(fetch_data=fetch_data))
    engine, session = make_session()
    row = _make_analyzed_upload(session, "f3" * 32)
    try:
        result = meta_analysis_service.analyze_meta_upload(session, row.id)

        assert [view for view, _ in calls] == ["v_ask_dispatch_pallet", "dblog_dispatch_pallet_log"]
        # Samma filter till bada vyerna, med numeriskt varde.
        assert calls[0][1] == [{"id": "pick_pall_num", "value": 8473877, "operator": "EQ"}]
        assert calls[1][1] == calls[0][1]
        assert result.analysis_status == "analyzed"
        assert result.order_number == "PR100485115"
        assert result.shipment_number == "GG-404-260610-PR100485115-0"
        assert result.username == "LOTS01"
        assert result.customer_name == "10231"  # arkivet saknar custom_desc; kundnr anvands
        assert result.uncertainty_notes is None
    finally:
        session.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_dispatch_lookup_unexpected_error_becomes_note_not_failure(monkeypatch):
    """Aven ett OVANTAT fel i ASK-uppslaget (inte bara kanda klientfel) ska
    bli en anteckning pa raden — analysen far aldrig kranga pa uppslaget."""
    monkeypatch.setattr(settings, "GEMINI_API_KEY", "gemini-key")
    monkeypatch.setattr(
        meta_analysis_service,
        "_call_meta_analysis_provider",
        lambda upload: {"pallet_id": "BOX-001", "deviations": ["Pall lutar"]},
    )
    monkeypatch.setattr(meta_analysis_service, "workflow_api_configured", lambda: True)

    def _boom():
        raise RuntimeError("ovantat uppslagsfel")

    monkeypatch.setattr(meta_analysis_service, "_api_client", _boom)
    engine, session = make_session()
    row = _make_analyzed_upload(session, "e2" * 32)
    try:
        result = meta_analysis_service.analyze_meta_upload(session, row.id)

        assert result.analysis_status == "manual_review"
        assert "kunde inte hamtas" in (result.uncertainty_notes or "")
        assert "BOX-001" in (result.uncertainty_notes or "")
        assert result.analysis_error is None
    finally:
        session.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_analyze_meta_upload_scrubs_api_key_from_analysis_error(monkeypatch):
    monkeypatch.setattr(settings, "GEMINI_API_KEY", "super-hemlig-nyckel")

    def _boom(upload):
        raise ValueError(
            "unknown url type: '#{GEMINI_API_BASE_URL}/upload/v1beta/files?key=super-hemlig-nyckel'"
        )

    monkeypatch.setattr(meta_analysis_service, "_call_meta_analysis_provider", _boom)
    engine, session = make_session()
    row = MetaMediaUpload(
        batch_id="batch-scrub",
        original_filename="voice.mov",
        stored_filename="20260531_120102_123456Z_01.mov",
        content_type="video/quicktime",
        media_type="video",
        size_bytes=11,
        content_hash="f" * 64,
        storage_backend="filesystem",
        storage_key=store_bytes(b"scrub-video", ".mov"),
        status="pending_analysis",
        source="public_upload",
    )
    session.add(row)
    session.commit()
    session.refresh(row)
    try:
        result = meta_analysis_service.analyze_meta_upload(session, row.id)

        assert result.analysis_status == "analysis_failed"
        error_text = result.analysis_error or ""
        assert "super-hemlig-nyckel" not in error_text
        assert "ValueError" in error_text
    finally:
        session.close()
        Base.metadata.drop_all(engine)
        engine.dispose()
